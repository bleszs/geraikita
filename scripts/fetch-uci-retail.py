"""Download original UCI workbook, lock content, derive country event sums."""
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import urllib.request
import uuid
import zipfile
from openpyxl import load_workbook
from uci_retail import aggregate_events, select_countries, calendar_dates, derived_series, country_key

ROOT = Path(__file__).resolve().parents[1]
PROTOCOL_PATH = ROOT / 'benchmarks/uci-online-retail/protocol.json'
BASE = ROOT / 'datasets/public/uci-online-retail'


def checksum(path, algorithm='sha256'):
    result = hashlib.new(algorithm)
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024*1024), b''):
            result.update(chunk)
    return result.hexdigest()


def write_json(path, data, exclusive=True):
    with path.open('x' if exclusive else 'w', encoding='utf-8', newline='\n') as stream:
        json.dump(data, stream, indent=2, allow_nan=False)
        stream.write('\n')


def main():
    protocol = json.loads(PROTOCOL_PATH.read_text(encoding='utf-8'))
    raw = BASE / 'raw'
    raw.mkdir(parents=True, exist_ok=True)
    archive, workbook = raw / 'online-retail.zip', raw / 'Online Retail.xlsx'
    lock_path = ROOT / 'benchmarks/uci-online-retail/source-lock.json'
    if not archive.exists():
        print('Downloading original UCI archive (about 23 MB).', flush=True)
        partial = raw / 'online-retail.zip.part'
        with urllib.request.urlopen(protocol['source']['downloadUrl'], timeout=60) as response, partial.open('wb') as out:
            size = 0
            while chunk := response.read(1024*1024):
                size += len(chunk)
                if size > 50_000_000:
                    raise ValueError('Unexpected archive size.')
                out.write(chunk)
        os.replace(partial, archive)
    with zipfile.ZipFile(archive) as source:
        info = source.getinfo('Online Retail.xlsx')
        if info.file_size > 50_000_000:
            raise ValueError('Unexpected workbook size.')
        payload = source.read(info)
    if workbook.exists():
        if checksum(workbook) != hashlib.sha256(payload).hexdigest():
            raise ValueError('Existing workbook differs from the downloaded archive.')
    else:
        workbook.write_bytes(payload)
    computed = {'archiveSha256': checksum(archive), 'workbookSha256': checksum(workbook),
                'workbookContentSha1': checksum(workbook, 'sha1'), 'archiveBytes': archive.stat().st_size, 'workbookBytes': workbook.stat().st_size}
    if lock_path.exists():
        lock = json.loads(lock_path.read_text(encoding='utf-8'))
        if lock['files'] != computed or lock['downloadUrl'] != protocol['source']['downloadUrl']:
            raise ValueError('Source content differs from frozen lock; stop without replacing it.')
    else:
        lock = {'lockedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'downloadUrl': protocol['source']['downloadUrl'], 'files': computed,
                'policy': 'Locally computed before parsing records. Not a publisher-supplied checksum. SHA1 is a content revision, not a Git commit.'}
        write_json(lock_path, lock)
    print('Source content pinned before parsing: ' + computed['workbookSha256'], flush=True)
    document = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if len(document.worksheets) != 1:
            raise ValueError('Unexpected number of source worksheets.')
        rows = iter(document.worksheets[0].iter_rows(values_only=True))
        header = next(rows)
        if list(header) != ['InvoiceNo','StockCode','Description','Quantity','InvoiceDate','UnitPrice','CustomerID','Country']:
            raise ValueError('Unexpected publisher columns.')
        quantities, event_counts, accounting = aggregate_events(rows, protocol)
    finally:
        document.close()
    if accounting['sourceRows'] != protocol['source']['expectedRows']:
        raise ValueError('Unexpected source row count; no replacement/fallback.')
    selected, eligibility = select_countries(quantities, protocol)
    development_dates = calendar_dates(protocol['calendar']['from'], protocol['calendar']['developmentThrough'])
    external_dates = calendar_dates(protocol['calendar']['externalFrom'], protocol['calendar']['externalThrough'])
    preparation = BASE / 'processed' / (dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H-%M-%SZ') + '-' + uuid.uuid4().hex[:8])
    preparation.mkdir(parents=True, exist_ok=False)
    write_json(preparation / 'selection.json', {'selected': selected, 'eligibility': eligibility, 'basis': 'Development interval only, before fitting or external scoring.'})
    entries = []
    for item in selected:
        country, scale = item['country'], item['earlyScale84']
        key = country_key(country)
        development = derived_series(country, development_dates, quantities, event_counts, scale)
        external = derived_series(country, external_dates, quantities, event_counts, scale)
        scope = {'workspaceId': 'public-uci-online-retail', 'businessId': 'country-'+key, 'productId': 'gross-product-units', 'unit': 'normalized_sales'}
        dataset = {'schemaVersion': 1, 'scope': scope, 'provenance': {
            'source': 'public_dataset', 'reference': protocol['source']['citation'], 'permissionConfirmed': True, 'completeDailyRecords': True,
            'usage': 'benchmark_only', 'sourceUrl': protocol['source']['url'], 'license': protocol['source']['license'],
            'sourceRevision': computed['workbookContentSha1'], 'sourceRevisionType': 'workbook_content_sha1_not_git', 'rawSha256': computed['workbookSha256'],
            'targetDescription': protocol['target'], 'aggregationPolicy': protocol['aggregation'], 'permissionBasis': 'Publisher CC BY 4.0 license, not private account consent.',
            'completenessBasis': 'Complete calendar aggregation of published event log, not verified complete real-world demand.',
            'normalization': {'scale': scale, 'from': development_dates[0], 'through': development_dates[83], 'policy': protocol['normalization']},
        }, 'observations': [{'date': r['date'], 'quantity': r['quantity']} for r in development]}
        train_file, external_file = key+'.train.json', key+'.external.json'
        write_json(preparation / train_file, dataset)
        write_json(preparation / external_file, external)
        write_json(preparation / (key+'.development-ledger.json'), development)
        entries.append({'country': country, 'key': key, 'scope': scope, 'scale': scale, 'trainFile': train_file, 'externalFile': external_file,
                        'trainSha256': checksum(preparation / train_file), 'externalSha256': checksum(preparation / external_file),
                        'developmentEmptyEventDates': sum(r['sourceEventRows'] == 0 for r in development)})
    manifest = {'protocol': protocol, 'protocolFileSha256': checksum(PROTOCOL_PATH), 'sourceLock': lock, 'accounting': accounting,
                'preparationDirectory': preparation.relative_to(ROOT).as_posix(), 'developmentDaysPerSeries': len(development_dates),
                'externalDaysPerSeries': len(external_dates), 'selection': eligibility, 'series': entries}
    write_json(preparation / 'manifest.json', manifest)
    pointer = BASE / 'processed/index.json'
    temporary = pointer.with_suffix('.tmp')
    write_json(temporary, manifest, exclusive=False)
    os.replace(temporary, pointer)
    print(json.dumps({'preparation': str(preparation), 'countriesSelected': len(entries), 'accounting': accounting, 'developmentDays': len(development_dates), 'externalDays': len(external_dates)}), flush=True)


if __name__ == '__main__':
    main()
