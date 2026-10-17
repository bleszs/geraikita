"""Source-only boosting fit, then fixed-weight transfer to new UCI groups."""
import os
os.environ['OMP_NUM_THREADS'] = '2'
os.environ['OPENBLAS_NUM_THREADS'] = '2'
os.environ['MKL_NUM_THREADS'] = '2'
import csv
import datetime as dt
import importlib.util
import json
from pathlib import Path
import platform
import uuid
import duckdb
import joblib
import numpy as np
import sklearn
from openpyxl import load_workbook
from threadpoolctl import threadpool_limits
from pooled_forecast import predictions, score
from sales_features import make_rows, feature_names, select_target_groups, METHODS
from uci_retail import aggregate_events, select_countries, calendar_dates, derived_series

ROOT = Path(__file__).resolve().parents[1]
PROTOCOL_PATH = ROOT / 'benchmarks/boosting-transfer/protocol.json'
spec = importlib.util.spec_from_file_location('source_benchmark', ROOT / 'scripts/benchmark-pooled.py')
source_benchmark = importlib.util.module_from_spec(spec)
spec.loader.exec_module(source_benchmark)
sha = source_benchmark.sha


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def write_json(path, data):
    with path.open('w', encoding='utf-8', newline='\n') as stream:
        json.dump(data, stream, indent=2, allow_nan=False)
        stream.write('\n')


def load_source(protocol):
    manifest_path = ROOT / protocol['source']['manifest']
    manifest = read_json(manifest_path)
    original = read_json(ROOT / protocol['source']['protocol'])
    if manifest['protocol'] != original or original['revision'] != protocol['source']['revision'] or len(manifest['series']) != 500:
        raise ValueError('Source manifest/protocol mismatch.')
    raw = ROOT / 'datasets/public/freshretailnet/raw/train.parquet'
    if sha(raw) != original['files']['train.parquet']['sha256']:
        raise ValueError('Source training checksum mismatch.')
    db = duckdb.connect()
    try:
        db.execute('SET threads=2')
        db.read_parquet(str(raw)).project('city_id,store_id,product_id,dt,sale_amount,stock_hour6_22_cnt').create_view('training')
        records = source_benchmark.read_series(db, [tuple(e['key']) for e in manifest['series']])
        for entry in manifest['series']:
            key = ':'.join(map(str, entry['key']))
            if records[key] != entry['trainStockouts']:
                raise ValueError('Source derived records differ from raw publisher input.')
    finally:
        db.close()
    # Discard stockout metadata before feature construction; target domain has none.
    return {key: [{'date': row['date'], 'quantity': row['quantity']} for row in series] for key, series in records.items()}


def validate_and_fit(protocol, development):
    candidates = []
    fitting = protocol['fitting']
    for variant in fitting['variants']:
        train = make_rows(development, variant, range(14, 76, 2))
        validation = make_rows(development, variant, range(76, 90, 2))
        for leaves in fitting['max_leaf_nodes']:
            model = source_benchmark.fit(train, leaves, fitting)
            for weight in fitting['residualWeights']:
                metrics = score([r['actual'] for r in validation], predictions(model, validation, weight))
                candidates.append({'variant': variant, 'maxLeafNodes': leaves, 'residualWeight': weight, **metrics})
            print(f'Source validation complete: {variant}, {leaves} leaves', flush=True)
    selected = min(candidates, key=lambda c: (c['mae'], c['residualWeight'], c['maxLeafNodes'], c['variant'] != 'lags'))
    all_rows = make_rows(development, selected['variant'], range(14, 90, 2))
    model = source_benchmark.fit(all_rows, selected['maxLeafNodes'], fitting)
    return model, selected, candidates, len(all_rows)


def load_targets(protocol):
    original = read_json(ROOT / protocol['target']['protocol'])
    lock = read_json(ROOT / protocol['target']['sourceLock'])
    raw = ROOT / 'datasets/public/uci-online-retail/raw/Online Retail.xlsx'
    if sha(raw) != lock['files']['workbookSha256']:
        raise ValueError('Target workbook checksum mismatch.')
    document = load_workbook(raw, read_only=True, data_only=True)
    try:
        if len(document.worksheets) != 1:
            raise ValueError('Unexpected target sheets.')
        events = iter(document.worksheets[0].iter_rows(values_only=True))
        if list(next(events)) != ['InvoiceNo','StockCode','Description','Quantity','InvoiceDate','UnitPrice','CustomerID','Country']:
            raise ValueError('Unexpected target columns.')
        quantities, counts, accounting = aggregate_events(events, original)
    finally:
        document.close()
    if accounting['sourceRows'] != original['source']['expectedRows']:
        raise ValueError('Target source row count changed.')
    selection_protocol = {**original, 'selection': {**original['selection'],
        'minimumActiveDates': protocol['target']['minimumActiveDates'], 'minimumUnits': protocol['target']['minimumUnits']}}
    _all_eligible, eligibility = select_countries(quantities, selection_protocol)
    selected = select_target_groups(eligibility, protocol['target']['excludedPreviouslyScoredCountries'])
    dates = calendar_dates(original['calendar']['from'], original['calendar']['externalThrough'])
    groups = {}
    for item in selected:
        # scale 1: raw invoiced quantity, no learned target-domain scaler/fitting.
        groups[item['country']] = derived_series(item['country'], dates, quantities, counts, 1)
    return groups, {'selected': selected, 'developmentEligibility': eligibility, 'excludedPreviouslyScoredCountries': protocol['target']['excludedPreviouslyScoredCountries'],
                    'accounting': accounting, 'sourceLock': lock, 'calendar': original['calendar']}


def summarize(records):
    return {method: score([r['actual'] for r in records], [r[method] for r in records]) for method in METHODS}


def evaluate(model, selected, groups):
    rows = make_rows(groups, selected['variant'], range(334, 372, 2))
    estimates = predictions(model, rows, selected['residualWeight'])
    return [{**{k: v for k, v in row.items() if k not in ['x', 'scale', 'base', 'residual']}, 'boosting': float(value)} for row, value in zip(rows, estimates)]


def target_summary(records, groups):
    pooled = summarize(records)
    countries = {country: summarize([r for r in records if r['key'] == country]) for country in groups}
    macro = {}
    for method in METHODS:
        defined = [scores[method]['wape'] for scores in countries.values() if scores[method]['wape'] is not None]
        macro[method] = {'meanCountryWape': sum(defined)/len(defined) if defined else None, 'definedCountries': len(defined), 'undefinedCountries': len(countries)-len(defined)}
    total = sum(r['actual'] for r in records)
    return {'targets': len(records), 'countries': len(groups), 'pooledRawUnitMetrics': pooled, 'primaryMacroCountry': macro, 'perCountry': countries,
            'targetVolumeShare': {country: sum(r['actual'] for r in records if r['key'] == country)/total if total else None for country in groups},
            'byOrigin': {date: summarize([r for r in records if r['date'] == date]) for date in sorted({r['date'] for r in records})},
            'countriesBeatingMA7': sum(scores['boosting']['mae'] < scores['ma7']['mae'] for scores in countries.values()),
            'countriesBeatingWeekly': sum(scores['boosting']['mae'] < scores['weekly']['mae'] for scores in countries.values())}


def main():
    protocol = read_json(PROTOCOL_PATH)
    run_id = dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H-%M-%SZ')+'-'+uuid.uuid4().hex[:8]
    directory = ROOT / 'models/benchmarks/boosting-transfer' / run_id
    directory.mkdir(parents=True, exist_ok=False)
    source_files = ['benchmarks/boosting-transfer/protocol.json', 'scripts/sales_features.py', 'scripts/benchmark-transfer.py',
                    'scripts/benchmark-pooled.py', 'scripts/pooled_forecast.py', 'scripts/uci_retail.py',
                    protocol['source']['protocol'], protocol['target']['protocol'], protocol['target']['sourceLock']]
    write_json(directory / 'frozen-configuration.json', {'frozenAt': dt.datetime.now(dt.timezone.utc).isoformat(),
               'sourceSha256': {file: sha(ROOT / file) for file in source_files}, 'priorKnownResults': protocol['priorKnowledge']})
    write_json(directory / 'protocol.json', protocol)
    development = load_source(protocol)
    model, selected, candidates, target_count = validate_and_fit(protocol, development)
    write_json(directory / 'validation.json', {'candidates': candidates, 'selected': selected, 'sourceOnly': True})
    artifact = directory / 'sales-only-model.joblib'
    joblib.dump(model, artifact, compress=3)
    metadata = {'modelKind': 'gerai-sales-only-residual-hgb-transfer-v1', 'usage': 'public_benchmark_only', 'trainingDataset': protocol['source'],
                'sourceSeries': len(development), 'sourceTrainingTargets': target_count, 'sourceTrainingThrough': '2024-06-25',
                'features': feature_names(selected['variant']), 'selected': selected, 'fitting': protocol['fitting'],
                'artifactSha256': sha(artifact), 'protocolSha256': sha(PROTOCOL_PATH), 'targetFittingPerformed': False,
                'frozenBeforeTargetParsingAt': dt.datetime.now(dt.timezone.utc).isoformat(),
                'environment': {'python': platform.python_version(), 'numpy': np.__version__, 'sklearn': sklearn.__version__, 'joblib': joblib.__version__}}
    write_json(directory / 'model-metadata.json', metadata)
    print('Model and configuration frozen before target workbook parsing: '+json.dumps(selected), flush=True)
    groups, selection = load_targets(protocol)
    write_json(directory / 'target-selection.json', selection)
    write_json(directory / 'target-records.json', groups)
    records = evaluate(model, selected, groups)
    if records != evaluate(joblib.load(artifact), selected, groups) or sha(artifact) != metadata['artifactSha256']:
        raise ValueError('Reload changed predictions or learned artifact.')
    write_json(directory / 'predictions.json', records)
    with (directory / 'predictions.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    result = target_summary(records, groups)
    report = {'schemaVersion': 1, 'generatedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'runId': run_id,
              'directory': directory.relative_to(ROOT).as_posix(), 'protocolVersion': protocol['protocolVersion'], 'model': metadata,
              'validation': {'selected': selected, 'candidates': candidates}, 'targetSelection': selection, 'external': result,
              'checks': {'targetFitOrTuning': False, 'previouslyScoredGroupOverlap': len(set(groups) & set(protocol['target']['excludedPreviouslyScoredCountries'])),
                         'savedModelReloadPredictionsIdentical': True, 'artifactUnchangedAfterTargetEvaluation': True, 'failedOrDroppedCountries': [], 'userDatabaseAccess': False},
              'limitations': protocol['limitations'], 'priorKnowledge': protocol['priorKnowledge']}
    write_json(directory / 'report.json', report)
    evaluation = ROOT / 'docs/evaluation'
    write_json(evaluation / 'boosting-transfer-results.json', report)
    write_json(directory.parent / 'latest.json', {'runId': run_id, 'directory': report['directory']})
    pct = lambda value: 'null' if value is None else f'{100*value:.2f}%'
    primary = '\n'.join(f"| {name} | {pct(scores['meanCountryWape'])} | {scores['undefinedCountries']} |" for name, scores in result['primaryMacroCountry'].items())
    secondary = '\n'.join(f"| {name} | {scores['mae']:.6f} | {pct(scores['wape'])} |" for name, scores in result['pooledRawUnitMetrics'].items())
    markdown = f'''# Transfer boosting sales-only ke kelompok UCI baru

Run `{run_id}`. Satu model baru dilatih hanya pada {len(development)} seri / {target_count:,} target dua hari FreshRetailNet, melalui 25 Juni 2024. Varian hanya memakai histori penjualan dan kalender; fitur stockout dan log level absolut dihapus sebelum eksperimen. Model v2 asli tetap tersimpan dan tidak dipaksakan menerima data stockout yang tidak tersedia.

Model, parameter dan checksum dibekukan sebelum workbook UCI diparsing pada run ini. Tidak ada fitting/tuning dengan UCI. Konteks 14 hari sebelumnya dari UCI dipakai untuk inferensi dengan bobot tetap. Test berisi **{len(groups)} negara dan {len(records)} target dua hari**, 1 November–8 Desember 2011. Negara terpilih: {', '.join(groups)}. Empat negara yang pernah dinilai pada replikasi Ridge dikeluarkan; tidak ada kelompok terpilih yang dibuang setelah scoring.

## Metrik utama: rata-rata WAPE per negara

| Metode | Mean country WAPE | Negara WAPE undefined |
|---|---:|---:|
{primary}

## Metrik tambahan: pooled unit invoice asli

| Metode | MAE / dua hari | WAPE |
|---|---:|---:|
{secondary}

Menang terhadap MA7: {result['countriesBeatingMA7']}/{len(groups)} negara; terhadap pola mingguan: {result['countriesBeatingWeekly']}/{len(groups)}. Konfigurasi dipilih pada validasi sumber saja: `{selected['variant']}`, {selected['maxLeafNodes']} leaf, bobot residual {selected['residualWeight']}. Bobot nol berarti keluaran sama dengan MA7 dan tidak membuktikan manfaat koreksi model.

## Batas klaim

Kalender UCI dan hasil empat negara sebelumnya sudah diketahui. Ini pengujian pada kelompok yang sebelumnya belum dinilai, bukan test yang sepenuhnya buta terhadap kondisi kalender. Agregat negara berasal dari satu peritel hadiah/grosir dan tidak mewakili usaha independen atau UMKM Indonesia. Target ialah unit invoice positif terpublikasi; SUM event kosong adalah nol menurut definisi target, bukan imputasi permintaan tak diketahui. Tidak ada fitur stockout buatan, konversi resep atau aktivasi model publik ke akun.

Model ini varian baru dengan bobot yang dipelajari dari FreshRetailNet dan ditransfer tetap; hasil tidak menguji bobot v2 asli atau dibandingkan langsung dengan angka sampel sebelumnya. Hasil kini diketahui, pengulangan hanya reproduksi. Qwen tetap pretrained. Seluruh negara, metrik, label, prediksi, source hash dan artefak disimpan.

[Protokol](../../benchmarks/boosting-transfer/protocol.json), [hasil lengkap](boosting-transfer-results.json). Artefak: `{report['directory']}`. Sumber/lisensi: [FreshRetailNet](https://huggingface.co/datasets/Dingdong-Inc/FreshRetailNet-50K) dan [UCI Online Retail](https://archive.ics.uci.edu/dataset/352/online+retail), CC BY 4.0.
'''
    (evaluation / 'BOOSTING_TRANSFER_RESULTS.md').write_text(markdown, encoding='utf-8')
    print(json.dumps({'runId': run_id, 'selected': selected, 'countries': len(groups), 'targets': len(records), 'primary': result['primaryMacroCountry'], 'pooled': result['pooledRawUnitMetrics']}), flush=True)


if __name__ == '__main__':
    with threadpool_limits(limits=2):
        main()
